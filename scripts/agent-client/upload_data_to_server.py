#!/usr/bin/env python3
"""
Upload 27 TDPs and 133 Subjects directly to QLMT Server PC via /api/import-data
"""

import sys
import os
import json
import math
import re
import urllib.request
import urllib.error
import ssl
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

ENDPOINT = "https://caplienchieu.dpdns.org/api/import-data"
STATUS_ENDPOINT = "https://caplienchieu.dpdns.org/api/agent/status"
IMPORT_SECRET = "b70ec0d7cfcafe5bfe549ca2e313bdea04e0813e35f16c59ed95c41e33cf8632"
AGENT_SECRET = "00fee6a637c79e5961f83ad1271a57edbd13b55ddb16f14c8f959dc9d8a83c5d"

PRESET_COLORS = [
    "#3b82f6", "#10b981", "#6366f1", "#8b5cf6",
    "#ec4899", "#f43f5e", "#06b6d4", "#14b8a6",
    "#f59e0b", "#10b981"
]


def format_tdp_name(raw_name: str) -> str:
    name = raw_name.strip()
    if re.match(r"^tổ\s+dân\s+phố\s+", name, re.IGNORECASE):
        name = re.sub(r"^tổ\s+dân\s+phố\s+", "Tổ dân phố ", name, flags=re.IGNORECASE)
    elif re.match(r"^tổ\s+", name, re.IGNORECASE):
        name = re.sub(r"^tổ\s+", "Tổ dân phố ", name, flags=re.IGNORECASE)
    else:
        name = "Tổ dân phố " + name

    words = name.split(" ")
    res = []
    for idx, word in enumerate(words):
        if word.upper() == "CC":
            res.append("CC")
        elif idx == 0 and re.match(r"^tổ$", word, re.IGNORECASE):
            res.append("Tổ")
        elif idx == 1 and re.match(r"^dân$", word, re.IGNORECASE):
            res.append("Dân")
        elif idx == 2 and re.match(r"^phố$", word, re.IGNORECASE):
            res.append("Phố")
        elif "-" in word:
            res.append("-".join(w.capitalize() for w in word.split("-")))
        else:
            res.append(word.capitalize())
    return " ".join(res)


def compute_polygon_centroid_and_area(coordinates):
    # Flatten outer ring if nested
    ring = coordinates[0] if coordinates and isinstance(coordinates[0][0], (list, tuple)) else coordinates
    if not ring or len(ring) < 3:
        return [16.075, 108.150], 10000

    # Shoelace formula for centroid and area
    n = len(ring)
    signed_area = 0.0
    cx = 0.0
    cy = 0.0

    for i in range(n - 1):
        x0, y0 = ring[i][0], ring[i][1]
        x1, y1 = ring[i + 1][0], ring[i + 1][1]
        a = x0 * y1 - x1 * y0
        signed_area += a
        cx += (x0 + x1) * a
        cy += (y0 + y1) * a

    if abs(signed_area) > 1e-12:
        signed_area *= 0.5
        cx = cx / (6.0 * signed_area)
        cy = cy / (6.0 * signed_area)
    else:
        cx = sum(p[0] for p in ring) / n
        cy = sum(p[1] for p in ring) / n

    # Area in sq meters (approximate at latitude 16 deg)
    # 1 deg lat ~ 110,649 m, 1 deg lng ~ 106,944 m
    area_sqm = abs(signed_area) * (110649.0 * 106944.0)
    return [round(cy, 6), round(cx, 6)], int(round(area_sqm))


def prepare_tdps():
    geojson_path = "data/final_tdps_2026.geojson"
    with open(geojson_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    features = data.get("features", [])
    tdp_list = []

    for i, f in enumerate(features):
        props = f.get("properties", {})
        raw_name = props.get("name", f"Tổ {i + 1}")
        formatted_name = format_tdp_name(raw_name)

        geom = f.get("geometry", {})
        coords = geom.get("coordinates", [])

        center, area_sqm = compute_polygon_centroid_and_area(coords)

        # Build clean FeatureCollection for this single TDP
        clean_feature = {
            "type": "Feature",
            "properties": {
                "name": formatted_name,
                "households": props.get("households", 0),
                "population": props.get("population", 0),
                "boundary_info": props.get("boundary_info", ""),
                "police_name": props.get("police_name", ""),
                "police_phone": props.get("police_phone", ""),
            },
            "geometry": geom,
        }

        tdp_doc = {
            "name": formatted_name,
            "households": props.get("households", 0),
            "population": props.get("population", 0),
            "area_sqm": area_sqm or int(props.get("area", 0) * 10000),
            "risk_status": props.get("risk_status", "green"),
            "color": props.get("color", PRESET_COLORS[i % len(PRESET_COLORS)]),
            "center": center,  # [lat, lng]
            "geojson": {
                "type": "FeatureCollection",
                "features": [clean_feature],
            },
            "boundary_info": props.get("boundary_info", ""),
            "police_name": props.get("police_name", ""),
            "police_phone": props.get("police_phone", ""),
            "leader_name": props.get("leader_name", ""),
            "leader_phone": props.get("leader_phone", ""),
        }
        tdp_list.append(tdp_doc)

    return tdp_list


# Named Area Mapping for matching address to TDP
NAMED_AREA_RULES = [
    (r"vân\s*dương\s*1|van\s*duong\s*1", "Tổ dân phố Vân Dương 1"),
    (r"vân\s*dương\s*2|van\s*duong\s*2", "Tổ dân phố Vân Dương 2"),
    (r"quan\s*nam\s*1", "Tổ dân phố Quan Nam 1"),
    (r"quan\s*nam\s*2", "Tổ dân phố Quan Nam 2"),
    (r"quan\s*nam\s*3|quan\s*nam\s*6", "Tổ dân phố Quan Nam 3"),
    (r"quan\s*nam\s*4", "Tổ dân phố Quan Nam 4"),
    (r"quan\s*nam\s*5", "Tổ dân phố Quan Nam 5"),
    (r"hưởng\s*phước|huong\s*phuoc", "Tổ dân phố Hưởng Phước"),
    (r"tân\s*ninh|tan\s*ninh|hiền\s*tân|hiền\s*phước|tân\s*hiền", "Tổ dân phố Tân Hiền"),
    (r"trung\s*sơn|trung\s*son|hồng\s*phước", "Tổ dân phố Trung Sơn - Hồng Phước"),
    (r"chung\s*cư|17\s*-\s*cc|17-cc", "Tổ dân phố Chung Cư 1"),
    (r"quang\s*thành\s*1\b", "Tổ dân phố Quang Thành 1"),
    (r"quang\s*thành\s*2\b", "Tổ dân phố Quang Thành 2"),
    (r"quang\s*thành\s*3\b", "Tổ dân phố Quang Thành 3"),
    (r"quang\s*thành\s*4\b", "Tổ dân phố Quang Thành 4"),
    (r"quang\s*thành\s*5\b", "Tổ dân phố Quang Thành 5"),
    (r"quang\s*thành\s*6\b", "Tổ dân phố Quang Thành 6"),
    (r"quang\s*thành\s*7\b", "Tổ dân phố Quang Thành 7"),
    (r"quang\s*thành\s*8\b", "Tổ dân phố Quang Thành 8"),
    (r"quang\s*thành\s*9\b", "Tổ dân phố Quang Thành 9"),
    (r"hòa\s*mỹ\s*1|hoa\s*my\s*1", "Tổ dân phố Hòa Mỹ 1"),
    (r"hòa\s*mỹ\s*2|hoa\s*my\s*2", "Tổ dân phố Hòa Mỹ 2"),
    (r"hòa\s*mỹ\s*3|hoa\s*my\s*3", "Tổ dân phố Hòa Mỹ 3"),
    (r"hòa\s*mỹ\s*4|hoa\s*my\s*4", "Tổ dân phố Hòa Mỹ 4"),
    (r"hòa\s*mỹ\s*5|hoa\s*my\s*5", "Tổ dân phố Hòa Mỹ 5"),
    (r"hòa\s*mỹ\s*6|hoa\s*my\s*6", "Tổ dân phố Hòa Mỹ 6"),
    (r"hòa\s*mỹ\s*7|hoa\s*my\s*7", "Tổ dân phố Hòa Mỹ 7"),
]


def match_tdp_for_address(addr: str, tdp_names: list) -> str:
    if not addr:
        return tdp_names[0]
    for pat, target in NAMED_AREA_RULES:
        if re.search(pat, addr, re.IGNORECASE):
            return target
    # Check if number mentioned (e.g. tổ 10, tổ 20)
    m = re.search(r"tổ\s*(\d+)", addr, re.IGNORECASE)
    if m:
        t_num = int(m.group(1))
        idx = (t_num - 1) % len(tdp_names)
        return tdp_names[idx]
    return tdp_names[0]


def is_valid_name(name) -> bool:
    if not name or not isinstance(name, str):
        return False
    trimmed = name.strip()
    if len(trimmed) < 2:
        return False
    if (
        trimmed.startswith("(")
        or trimmed.startswith("STT")
        or "Họ và tên" in trimmed
        or "CÁN BỘ" in trimmed
        or "TRƯỞNG" in trimmed
        or "DANH SÁCH" in trimmed
        or "Biểu mẫu" in trimmed
        or "ĐƠN VỊ" in trimmed
        or trimmed.startswith("Lợi:")
        or trimmed.startswith("Hùng:")
        or trimmed.startswith("Quý:")
    ):
        return False
    return True


def format_date_val(val):
    if not val:
        return ""
    if isinstance(val, (int, float)):
        # Excel date offset from 1899-12-30
        try:
            import datetime
            base = datetime.datetime(1899, 12, 30)
            d = base + datetime.timedelta(days=val)
            return d.strftime("%d/%m/%Y")
        except Exception:
            return str(val)
    if hasattr(val, "strftime"):
        return val.strftime("%d/%m/%Y")
    return str(val).strip()


def extract_yob(dob_str: str) -> int:
    if not dob_str:
        return 0
    m = re.search(r"\b(19\d\d|20\d\d)\b", dob_str)
    return int(m.group(1)) if m else 0


def prepare_subjects(tdps):
    tdp_map = {t["name"]: t["center"] for t in tdps}
    tdp_names = list(tdp_map.keys())

    excel_path = "data/raw/DS ma túy Mẫu 1, Mẫu 2 bổ sung 27.01.2026.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    sheet_configs = [
        ("Mẫu 1", "Sử dụng"),
        ("Mẫu 2", "Nghiện"),
        ("Mẫu 3", "Sau cai"),
        ("Mẫu 4", "Khởi tố"),
        ("Mẫu 5", "Quản lý sau cai"),
        ("Mẫu 6", "Quản lý sau cai"),
        ("Thanh loại", "Thanh loại"),
    ]

    subjects = []
    tdp_cluster_count = {}

    for sheet_name, default_status in sheet_configs:
        if sheet_name not in wb.sheetnames:
            continue
        sheet = wb[sheet_name]

        for row in sheet.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            name_cand = row[1]
            if not is_valid_name(name_cand):
                continue

            full_name = str(name_cand).strip()
            raw_dob = row[2] if len(row) > 2 else ""
            dob_str = format_date_val(raw_dob)
            yob = extract_yob(dob_str)
            gender = str(row[3]).strip() if len(row) > 3 and row[3] else "Nam"
            id_card = str(row[4]).strip() if len(row) > 4 and row[4] else None
            addr_perm = str(row[5]).strip() if len(row) > 5 and row[5] else ""
            addr_curr = str(row[6]).strip() if len(row) > 6 and row[6] else addr_perm

            status = default_status
            if len(row) > 12 and row[12] and isinstance(row[12], str) and row[12].strip():
                status = row[12].strip()

            violation_histories = []
            col7 = format_date_val(row[7]) if len(row) > 7 else ""
            col8 = str(row[8]).strip() if len(row) > 8 and row[8] else ""
            col9 = str(row[9]).strip() if len(row) > 9 and row[9] else ""
            col10 = str(row[10]).strip() if len(row) > 10 and row[10] else ""
            col11 = str(row[11]).strip() if len(row) > 11 and row[11] else ""

            if sheet_name in ["Mẫu 1", "Thanh loại"]:
                date_val = col7
                decision_val = col8
                duration_val = col9 or col10
            else:
                date_val = col7
                decision_val = " - ".join(x for x in [col8, col9, col10] if x)
                duration_val = col11

            if date_val or decision_val or duration_val:
                violation_histories.append({
                    "action": status,
                    "date": date_val,
                    "decision_num_date": decision_val,
                    "duration": duration_val,
                })

            notes_items = []
            for col_idx in [12, 13, 14]:
                if len(row) > col_idx and row[col_idx]:
                    notes_items.append(str(row[col_idx]).strip())
            notes = " | ".join(notes_items)

            # Match TDP
            matched_tdp = match_tdp_for_address(addr_curr or addr_perm, tdp_names)
            center = tdp_map.get(matched_tdp, [16.075, 108.150])

            # Cluster offset for distinct map markers
            cluster_idx = tdp_cluster_count.get(matched_tdp, 0) + 1
            tdp_cluster_count[matched_tdp] = cluster_idx

            angle = (cluster_idx * 137.5) * (math.pi / 180.0)
            dist = 0.00015 * math.sqrt(cluster_idx)
            lat = round(center[0] + dist * math.sin(angle), 6)
            lng = round(center[1] + dist * math.cos(angle), 6)

            sub_doc = {
                "full_name": full_name,
                "dob": dob_str,
                "yob": yob,
                "gender": gender,
                "id_card": id_card,
                "ethnicity": "Kinh",
                "address_permanent": addr_perm,
                "address_current": addr_curr,
                "tdp": matched_tdp,
                "drug_types_used": ["MET", "OPI"],
                "status": status,
                "violation_histories": violation_histories,
                "notes": notes,
                "lat": lat,
                "lng": lng,
                "is_drug": 1,
                "approval_status": "Approved",
            }
            subjects.append(sub_doc)

    return subjects


def post_import(collection_name, data):
    payload = json.dumps({
        "collection": collection_name,
        "secret": IMPORT_SECRET,
        "data": data,
    }).encode("utf-8")

    req = urllib.request.Request(
        ENDPOINT,
        data=payload,
        headers={
            "Content-Type": "application/json",
            "User-Agent": "LaptopAgentUploader/1.0",
        },
        method="POST",
    )

    ctx = ssl.create_default_context()
    with urllib.request.urlopen(req, context=ctx, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def check_status():
    req = urllib.request.Request(
        STATUS_ENDPOINT,
        headers={"x-agent-secret": AGENT_SECRET},
        method="GET",
    )
    ctx = ssl.create_default_context()
    with urllib.request.urlopen(req, context=ctx, timeout=15) as resp:
        return json.loads(resp.read().decode("utf-8"))


def main():
    print("🚀 Đang chuẩn bị dữ liệu 27 TDP từ final_tdps_2026.geojson...")
    tdps = prepare_tdps()
    print(f"✅ Đã chuẩn bị {len(tdps)} TDP.")

    print("\n🚀 Đang chuẩn bị dữ liệu 133 Đối tượng từ file Excel...")
    subjects = prepare_subjects(tdps)
    print(f"✅ Đã chuẩn bị {len(subjects)} Đối tượng ma túy.")

    print(f"\n📤 1. Đang gửi {len(tdps)} TDP lên Server ({ENDPOINT})...")
    res_tdp = post_import("tdps", tdps)
    print("Kết quả import TDP:", json.dumps(res_tdp, ensure_ascii=False))

    print(f"\n📤 2. Đang gửi {len(subjects)} Đối tượng lên Server ({ENDPOINT})...")
    res_sub = post_import("subjects", subjects)
    print("Kết quả import Đối tượng:", json.dumps(res_sub, ensure_ascii=False))

    print("\n🔍 3. Kiểm tra lại trạng thái database trên máy chủ qua Agent Status...")
    status = check_status()
    print("Trạng thái mới trên Server:", json.dumps(status, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
